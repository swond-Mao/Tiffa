# -*- coding: utf-8 -*-
"""对码库 30 个 xlsx → SQLite 索引（流式读取，可断点重跑：文件级去重）"""
import openpyxl, os, re, sqlite3, sys, time, unicodedata

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB = os.path.join(BASE, 'data', 'duima_index.sqlite')
os.makedirs(os.path.dirname(DB), exist_ok=True)

def norm(s):
    if s is None:
        return ''
    s = unicodedata.normalize('NFKC', str(s))
    s = s.lower().replace('度', '°')
    s = re.sub(r'[\s\-_～~・·]+', '', s)
    s = s.replace('×', '*').replace('x', '*').replace('×', '*')
    return s.strip()

# 列下标（0 起）：3 医保耗材分类代码 4 一级 5 二级 6 三级 14 注册备案产品名称 15 单件产品名称 16 耗材企业 18 规格 19 型号 23 27位码
COLS = dict(code=3, c1=4, c2=5, c3=6, reg=14, item=15, ent=16, spec=18, model=19, code27=23)

conn = sqlite3.connect(DB)
conn.execute('PRAGMA journal_mode=WAL')
conn.execute('PRAGMA synchronous=NORMAL')
conn.execute('''CREATE TABLE IF NOT EXISTS sku(
  id INTEGER PRIMARY KEY, src TEXT, c1 TEXT, c3 TEXT,
  reg_name TEXT, item_name TEXT, ent TEXT,
  spec TEXT, model TEXT, code27 TEXT,
  n_model TEXT, n_spec TEXT, n_item TEXT, n_reg TEXT)''')
conn.execute('CREATE INDEX IF NOT EXISTS idx_n_model ON sku(n_model)')
conn.execute('CREATE INDEX IF NOT EXISTS idx_n_spec ON sku(n_spec)')
conn.execute('CREATE INDEX IF NOT EXISTS idx_n_item ON sku(n_item)')
conn.execute('CREATE TABLE IF NOT EXISTS done_files(name TEXT PRIMARY KEY, rows INTEGER)')
conn.commit()

files = sorted(f for f in os.listdir(BASE) if f.startswith('医保医用耗材代码_') and f.endswith('.xlsx'))
print(f'共 {len(files)} 个文件', flush=True)

t_all = time.time()
for f in files:
    t0 = time.time()
    if conn.execute('SELECT 1 FROM done_files WHERE name=?', (f,)).fetchone():
        print(f'[skip] {f}', flush=True)
        continue
    wb = openpyxl.load_workbook(os.path.join(BASE, f), read_only=True)
    ws = wb[wb.sheetnames[0]]
    batch, n = [], 0
    for i, r in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        code27 = r[COLS['code27']]
        if code27 is None or not str(code27).strip():
            continue
        row = (
            f, r[COLS['c1']] or '', r[COLS['c3']] or '',
            str(r[COLS['reg']] or ''), str(r[COLS['item']] or ''), str(r[COLS['ent']] or ''),
            str(r[COLS['spec']] or ''), str(r[COLS['model']] or ''), str(code27).strip(),
            norm(r[COLS['model']]), norm(r[COLS['spec']]), norm(r[COLS['item']]), norm(r[COLS['reg']]),
        )
        batch.append(row)
        n += 1
        if len(batch) >= 50000:
            conn.executemany('INSERT INTO sku VALUES(NULL,?,?,?,?,?,?,?,?,?,?,?,?,?)', batch)
            conn.commit()
            batch = []
    if batch:
        conn.executemany('INSERT INTO sku VALUES(NULL,?,?,?,?,?,?,?,?,?,?,?,?,?)', batch)
        conn.commit()
    wb.close()
    conn.execute('INSERT INTO done_files VALUES(?,?)', (f, n))
    conn.commit()
    print(f'[done] {f}  rows={n}  {time.time()-t0:.0f}s  (total {time.time()-t_all:.0f}s)', flush=True)

total = conn.execute('SELECT COUNT(*) FROM sku').fetchone()[0]
print(f'ALL DONE. sku total = {total}', flush=True)
conn.close()
