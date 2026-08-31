# -*- coding: utf-8 -*-
"""把 耗材分类总表_已对码.xlsx 的总表按使用年份拆成 近1年/近2年/近3年 三个 sheet
口径（与源表 最后使用年份 列一致）：
  近1年 = 最后使用年份 == 2026
  近2年 = 最后使用年份 in (2025, 2026)
  近3年 = 全表
"""
import os
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
F = os.path.join(BASE, '耗材分类总表_已对码.xlsx')

wb = openpyxl.load_workbook(F)
ws = wb['总表']
hdr = [c.value for c in ws[1]]
yi = hdr.index('最后使用年份')
rows = list(ws.iter_rows(values_only=True))
data = rows[1:]

def make_sheet(title, pred):
    if title in wb.sheetnames:
        del wb[title]
    st = wb.create_sheet(title)
    st.append(hdr)
    n = 0
    for r in data:
        if pred(r[yi]):
            st.append(list(r))
            n += 1
    print(f'{title}: {n} 行')

make_sheet('近1年', lambda y: str(y) == '2026')
make_sheet('近2年', lambda y: str(y) in ('2025', '2026'))
make_sheet('近3年', lambda y: True)
wb.save(F)
print('saved', F)
