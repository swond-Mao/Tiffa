# -*- coding: utf-8 -*-
"""把 对码统计 sheet 重写为按时间窗(近1年/近2年/近3年)的状态矩阵 + 覆盖率"""
import os
from collections import Counter
import openpyxl

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
F = os.path.join(BASE, '耗材分类总表_已对码.xlsx')

wb = openpyxl.load_workbook(F)
ws = wb['总表']
hdr = [c.value for c in ws[1]]
yi = hdr.index('最后使用年份')
ci = hdr.index('医保27位码')
si = hdr.index('对码状态')

g = {'近1年': Counter(), '近2年': Counter(), '近3年': Counter()}
coded = {'近1年': 0, '近2年': 0, '近3年': 0}
total = {'近1年': 0, '近2年': 0, '近3年': 0}
for r in ws.iter_rows(min_row=2, values_only=True):
    y = str(r[yi])
    rec = (r[si], bool(r[ci]))
    for name in ('近3年',):
        pass
    g['近3年'][rec[0]] += 1
    total['近3年'] += 1
    coded['近3年'] += 1 if rec[1] else 0
    if y in ('2025', '2026'):
        g['近2年'][rec[0]] += 1
        total['近2年'] += 1
        coded['近2年'] += 1 if rec[1] else 0
    if y == '2026':
        g['近1年'][rec[0]] += 1
        total['近1年'] += 1
        coded['近1年'] += 1 if rec[1] else 0

STATS = ['厂家型号唯一命中', '厂家型号取Top1', '型号唯一命中', '多候选取Top1',
         '名称匹配取Top1', '名称唯一命中', '厂家型号待人工', '多候选待人工',
         '未匹配', '非耗材(外配药品)']
AUTO = {'厂家型号唯一命中', '厂家型号取Top1', '型号唯一命中', '多候选取Top1', '名称匹配取Top1', '名称唯一命中'}
MANUAL = {'厂家型号待人工', '多候选待人工'}

if '对码统计' in wb.sheetnames:
    del wb['对码统计']
st = wb.create_sheet('对码统计')
st.append(['对码状态', f"近1年({total['近1年']})", f"近2年({total['近2年']})", f"近3年({total['近3年']})"])
for s in STATS:
    st.append([s] + [g[w][s] for w in ('近1年', '近2年', '近3年')])
st.append([])
st.append(['指标', '近1年', '近2年', '近3年'])
st.append(['已填27位码数'] + [coded[w] for w in ('近1年', '近2年', '近3年')])
st.append(['覆盖率(含待人工建议码)'] + [f"{coded[w]/total[w]:.1%}" for w in ('近1年', '近2年', '近3年')])
st.append(['可直接采纳数'] + [sum(g[w][s] for s in AUTO) for w in ('近1年', '近2年', '近3年')])
st.append(['可直接采纳占比'] + [f"{sum(g[w][s] for s in AUTO)/total[w]:.1%}" for w in ('近1年', '近2年', '近3年')])
st.append(['待人工确认数'] + [sum(g[w][s] for s in MANUAL) for w in ('近1年', '近2年', '近3年')])
st.append(['未匹配数'] + [g[w]['未匹配'] for w in ('近1年', '近2年', '近3年')])
for col, w in zip('ABCD', (24, 16, 16, 16)):
    st.column_dimensions[col].width = w
wb.save(F)
print('对码统计 sheet 已更新')
